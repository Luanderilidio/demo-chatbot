from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
    for border in list(s.element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(13)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(6)
doc.core_properties.title='Evolução e arquitetura do chatbot agronômico'
doc.core_properties.subject='RAG com LangChain LangGraph e API de fazenda'
doc.core_properties.author='Equipe do projeto'
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Relatório técnico • ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

def p(text): return doc.add_paragraph(text)
def h(text): return doc.add_heading(text,2)
def page(title):
    heading=doc.add_heading(title,1)
    heading.paragraph_format.page_break_before=True
def bullets(items):
    for text in items: doc.add_paragraph(text,style='List Bullet')
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for cell,text in zip(t.rows[0].cells,headers): cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit'); pr.append(ns)
        if i==0: pr.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'203C50' if i==0 else ('F0F4F6' if i%2 else 'FFFFFF')); cp.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
            cp.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3); para.paragraph_format.space_before=Pt(3)
                for run in para.runs:
                    run.font.size=Pt(10)
                    if i==0: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    spacer=doc.add_paragraph()
    spacer.paragraph_format.line_spacing=Pt(3)
    spacer.paragraph_format.space_after=Pt(4)
    return t

doc.add_paragraph('Evolução e arquitetura do chatbot agronômico',style='Title')
doc.add_paragraph('Implementação com RAG LangChain LangGraph e API de fazenda',style='Subtitle')
p('Relatório técnico • 13 de outubro de 2026 • Versão 1')
h('Objetivo e conclusão')
p('O projeto evoluiu de um protótipo de consulta a documentos para uma aplicação local que combina recuperação de conhecimento agronômico, medições meteorológicas simuladas e um plano fictício de plantio. A integração está funcional para demonstração e testes. Ainda não constitui um sistema validado de recomendação agronômica ou autorização de plantio.')
p('Este documento explica a implementação existente, a função de cada tecnologia, o processamento de tokens, os resultados observados e o trabalho necessário para transformar o protótipo em um serviço confiável. O público principal é a equipe de desenvolvimento e os profissionais que irão revisar os critérios agronômicos.')
h('O que existe e o que ainda é proposta')
table(['Situação','Componentes'],[
('Implementado','Interface de terminal Python; RAG documental; dois grafos LangGraph; modelos locais no Ollama; API JSON Server; plano e validações de dados simulados.'),
('Verificado em testes','Percursos do grafo, tratamento de falhas, associação entre cadastro e fazenda, comparação numérica e integração HTTP. Há exemplos com Ollama e Chroma reais.'),
('Não implementado','Aplicativo web ou móvel, backend público do chatbot, autenticação, memória persistente, sensores reais, previsão meteorológica e integração oficial de zoneamento.'),
('Ainda não demonstrado','Precisão agronômica em amostra representativa, confiabilidade de decisão em campo e desempenho sob múltiplos usuários.')],[3.3,13.7])
h('Como ler este relatório')
p('As seções 1 a 5 descrevem evolução, tecnologias, arquitetura e dados. As seções 6 a 8 explicam tokens, exemplos e validação. As seções 9 a 11 apresentam os dados reais, a arquitetura futura e os critérios de aceitação. A seção 12 reúne referências e comandos.')
p('LangChain e LangGraph são os componentes utilizados. Não há blockchain nesta arquitetura. O armazenamento da base vetorial também não modifica os pesos do modelo: RAG e treinamento são processos diferentes.')

page('1 Evolução do projeto')
h('Preparação do protótipo recebido')
p('A etapa inicial examinou os arquivos recebidos, as dependências Python e a execução local. Foram tratados problemas de ambiente e de abertura do índice vetorial. Nesta instalação, o banco ativo foi mantido em um caminho externo sem acentos; a cópia original permaneceu como backup. Isso é uma característica observada neste ambiente, não uma exigência universal do Chroma.')
p('A organização separou entradas de execução, implementação em src, configuração, documentos, scripts, estudos e testes. O ambiente utilizado passou a ser .venv312. A pasta do projeto posteriormente recebeu o nome chatbot.')
h('Orquestração com LangGraph')
p('O fluxo linear de buscar e responder recebeu avaliação de evidências, reformulação limitada e saída de esclarecimento. A pergunta original permanece no estado durante uma consulta. Há no máximo duas buscas documentais. Falhas de avaliação estruturada e falhas de infraestrutura são distinguidas da ausência de informação.')
h('Otimização de desempenho')
p('O diagnóstico identificou que a avaliação dos trechos consumia a maior parte da latência. A saída do avaliador foi encurtada, eliminando IDs duplicados e justificativas extensas. A janela de contexto passou de 8.192 para 4.096 tokens, e o modelo foi configurado para permanecer carregado por até 15 minutos. Uma regressão na seleção de fontes foi detectada durante os testes e corrigida antes de aceitar a mudança.')
h('Integração da fazenda simulada')
p('O JSON Server passou a disponibilizar medições e, depois, um plano de plantio com solo, cultivar e janela fictícia. O grafo externo escolhe documentos, API ou ambos. Consultas simples à API dispensam geração por LLM. Para uma faixa explícita de temperatura de semeadura, o Python realiza uma comparação verificável com os valores obtidos.')
h('Correções de consistência e operação')
bullets(['A umidade volumétrica do solo passou a ter uma única origem no plano do talhão; o campo duplicado das medições foi removido.',
'Planos ausentes, inválidos, antigos ou com restrições bloqueiam a análise combinada de plantio antes da geração.',
'A API passou a recarregar o JSON alterado e a rejeitar conteúdo inválido, sem reutilizar silenciosamente dados antigos.',
'A mudança do nome da pasta revelou vínculos absolutos quebrados nas dependências Node. A instalação foi refeita com nodeLinker hoisted.',
'Mensagens de erro foram ajustadas para distinguir servidor indisponível, recurso HTTP ausente e porta já ocupada.'])

page('2 Tecnologias utilizadas')
p('As versões abaixo foram verificadas no ambiente Python local para este relatório. Elas não substituem um arquivo de dependências reproduzível: requirements.txt ainda contém faixas abertas em várias bibliotecas.')
table(['Tecnologia','Versão ou modelo','Responsabilidade no projeto'],[
('Python','3.12 no ambiente .venv312','Executa o chatbot, valida dados e coordena chamadas.'),
('LangChain Core','1.6.10','Mensagens, Document, prompts e composição de operações.'),
('LangGraph','1.2.14','Nós, estado, transições condicionais e streaming de eventos.'),
('LangChain Ollama','1.1.0','Integra o Python aos modelos locais.'),
('Ollama e Llama','llama3.2:3b','Serviço de inferência e modelo gerador/avaliador; 3b indica aproximadamente três bilhões de parâmetros.'),
('Embeddings','nomic-embed-text','Representa texto em vetores para recuperação semântica.'),
('Chroma e integração','1.5.9 e langchain-chroma 1.1.0','Persiste vetores, textos e metadados; recupera trechos por similaridade.'),
('Pydantic','2.14.1','Valida formatos, tipos, limites e estruturas recebidas.'),
('pypdf e Rich','6.20.0 e 15.0.0','Extração textual de PDFs e interface de terminal.'),
('Node e JSON Server','24.19.0 e 0.17.4','Executam a API REST fictícia sobre api/db.json.'),
('pnpm e unittest','Gerenciador e biblioteca padrão','Instala dependências Node e executa testes Python.')],[3.1,4.2,9.7])
p('LangChain conecta as operações; LangGraph controla a ordem e as condições de execução. O LLM produz linguagem e classificações, mas não é o banco de dados. Chroma responde à busca documental; JSON Server expõe dados estruturados. A documentação oficial distingue essas responsabilidades [R1–R4].')
p('No protótipo, não estão em uso LangSmith, treinamento supervisionado do modelo, fine tuning, aplicação móvel, serviços de nuvem ou blockchain. Recursos disponíveis nas bibliotecas não devem ser confundidos com recursos efetivamente configurados.')

page('3 Arquitetura e fluxo de execução')
p('A aplicação atual é uma CLI. Ela não possui uma API própria de chat: chama diretamente AgriculturalRAG no processo Python. O HTTP existente atende os dados da fazenda, enquanto o Ollama atende inferência local. O banco vetorial é acessado pela integração Chroma.')

# Diagrama técnico legível dentro da página.
im=Image.new('RGB',(1500,590),'white'); draw=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24)
def box(x,y,w,ht,lines):
    draw.rounded_rectangle((x,y,x+w,y+ht),radius=9,fill='#EEF3F5',outline='#203C50',width=3)
    for i,line in enumerate(lines): draw.text((x+18,y+15+i*30),line,font=font,fill='black')
def arrow(a,b):
    draw.line((a,b),fill='#203C50',width=4)
    x,y=b; draw.polygon([(x,y),(x-9,y-16),(x+9,y-16)],fill='#203C50')
box(510,12,480,80,['Pergunta na interface de terminal'])
arrow((750,92),(750,135)); box(420,135,660,85,['Grafo externo LangGraph','Roteamento por regras e validação'])
for x in [240,750,1260]:
    draw.line(((750,220),(x,260)),fill='#203C50',width=3); arrow((x,260),(x,290))
box(30,290,420,100,['API de fazenda','HTTP e resposta direta'])
box(540,290,420,100,['Plano e documentos','Validação e comparação'])
box(1050,290,420,100,['RAG documental','Chroma e Ollama'])
for x in [240,750,1260]: draw.line(((x,390),(750,445)),fill='#203C50',width=3)
arrow((750,445),(750,480)); box(350,480,800,90,['Resposta com origem dos dados e limitações','ou esclarecimento ou interrupção por dados insuficientes'])
im.save(OUT/'arquitetura.png')
doc.add_picture(str(OUT/'arquitetura.png'),width=Cm(17))
p('Figura 1. Fluxo atual. A comparação determinística e a geração por modelo são caminhos distintos, escolhidos conforme a operação e o formato da evidência.')
h('Duas camadas de grafo')
p('assistant_graph.py roteia entre responder_api, documentos e combinar_fontes. O roteamento usa regras de texto em português; não é um classificador de intenção universal. A consulta simples a sensores ou ao plano pode terminar sem chamar o modelo.')
p('rag_graph.py implementa buscar, avaliar, reformular, responder, esclarecer e sem_evidencia. O avaliador retorna decisão estruturada e passagens. A aplicação valida a atribuição das passagens antes de aprovar fontes. Se a primeira busca for insuficiente, há uma reformulação; a segunda tentativa encerra o ciclo.')
h('Estado e memória')
p('O estado guarda pergunta original, consulta de busca, filtro, número de tentativas, documentos, avaliação e resultado. API e plano são obtidos para a consulta em curso. Não há histórico persistente, checkpoint ou memória entre mensagens; “e amanhã?” não recebe automaticamente o contexto da pergunta anterior.')

page('4 Preparação da base e recuperação documental')
h('Ingestão existente')
p('src/ingest.py seleciona arquivos pelo DOC_METADATA_MAP, extrai texto por página com pypdf e trata relatórios Markdown. O mapa contém seis PDFs e dois relatórios técnicos. Existe outro PDF no acervo que não está nesse mapa; presença na pasta não comprova que um documento foi indexado.')
p('O texto é limpo e dividido com RecursiveCharacterTextSplitter: chunk_size=750 e chunk_overlap=120. Nessa configuração, os tamanhos são medidos em caracteres, não em tokens. O embedding de cada trecho é calculado pelo nomic-embed-text e armazenado na coleção embrapa_conhecimento junto ao texto e aos metadados.')
table(['Dado persistido','Uso'],[
('Texto do trecho','Conteúdo fornecido ao avaliador e ao gerador.'),
('Vetor de embedding','Busca por proximidade semântica; não é uma probabilidade de verdade.'),
('source e titulo','Identificação do arquivo e nome da publicação.'),
('cultura','Filtro de recuperação por domínio/cultura.'),
('pagina e total_paginas','Localização no arquivo; pagina usa o índice do PDF mais um, que pode divergir da paginação impressa.')],[5,12])
h('Recuperação durante uma pergunta')
p('A pergunta é convertida em embedding; Chroma retorna os quatro candidatos configurados por k=4, respeitando o filtro quando ativo. O avaliador verifica se os trechos sustentam uma resposta. Até duas passagens podem compor a seleção compacta. Comparar por similaridade semântica aproxima assuntos; não garante que a pergunta foi respondida ou que o trecho contém todas as condições necessárias [R2, R4].')
h('Limitações que ainda precisam ser corrigidas')
p('A retomada da ingestão usa a contagem já existente para cortar a lista de trechos. Isso não detecta corretamente arquivos alterados, removidos ou reordenados. Além disso, o código pode exibir conclusão após falha persistente de um lote. O teste de diagnóstico existente reproduz essas falhas; ele não significa que elas foram resolvidas.')
p('Antes de atualizar o corpus em produção, são necessários IDs estáveis por documento/versão/trecho, hashes do conteúdo, atualização idempotente, remoção de versões antigas e confirmação dos lotes realmente gravados. Deve-se preservar a estrutura das perguntas e respostas e recuperar passagens vizinhas quando uma condição atravessar uma página.')

page('5 API e contratos de dados')
p('O JSON Server expõe REST em 127.0.0.1:3001. O cliente Python usa GET com timeout de três segundos por requisição e valida os dados. Não há cache meteorológico no cliente. O servidor recarrega alterações externas do JSON; conteúdo inválido produz HTTP 503. Esse comportamento foi testado com um banco temporário.')
table(['Recurso','Campos principais','Papel'],[
('/fazendas','id, nome, município, UF, fuso','Identifica a unidade consultada.'),
('/medicoes','fazendaId, talhaoId, observado_em, temperaturas, chuva, vento, umidade do ar','Última observação meteorológica simulada.'),
('/planos_plantio','cultura, cultivar, ciclo, solo, horário, janela e regras de teste','Cadastro do talhão e condição hídrica simulada.'),
('/health','service, version, simulado','Identifica o processo da API; não comprova validade agronômica.')],[4,8,5])
h('Uma origem para cada informação')
p('A umidade do solo fica somente em planos_plantio.umidade_volumetrica_pct. O antigo campo umidade_solo_pct das medições foi removido para impedir respostas contraditórias. Umidade volumétrica de 28% é uma observação distinta dos 68% de umidade relativa do ar. Chuva acumulada de 12 mm também não é uma medição da umidade do solo.')
p('Os vínculos entre fazenda, talhão, cultura, município e UF são conferidos antes de aceitar o plano. Mais de um plano aplicável ou medições de talhões diferentes não são resolvidos por escolha arbitrária. A demonstração ainda não possui uma interface completa de seleção de talhão.')
h('Atualidade e regras artificiais')
p('A observação deve ser do mesmo dia em Cuiabá, sem horário futuro e com idade de até seis horas. Esse limite é uma política do protótipo; não foi validado como frequência agronômica adequada. A observação do solo possui horário próprio. Datas de início e fim da janela são inclusivas.')
p('A cultivar SOJA-DEMO-001, seu ciclo de 115 dias, a janela de 01/10 a 30/11/2026 e o intervalo hídrico de 20–35% foram inventados para testar o software. A fonte da janela é SIMULACAO_NAO_ZARC. Nenhuma dessas informações representa validação oficial da cultivar ou consulta ao zoneamento.')
p('Executar seed.cjs substitui o JSON pelos valores padrão e atualiza horários. Isso cria novos dados fictícios; não coleta sensores. Alterar JSON ou executar seed não retreina o LLM e não requer nova ingestão no Chroma.')

page('6 Tokens e custo de processamento')
h('O que é um token')
p('Token é uma unidade de representação do texto usada pelo modelo. Pode corresponder a uma palavra, parte dela, pontuação ou combinação de caracteres. O tokenizador transforma essas unidades em IDs. Não existe equivalência fixa entre palavras, caracteres e tokens; a divisão depende do vocabulário do modelo e do idioma [R5].')
p('Por exemplo, “Qual é a temperatura do solo?” precisa ser tokenizado pelo modelo utilizado para obter a contagem real. Separar a frase por espaços não fornece essa contagem. O Llama e o modelo de embeddings podem empregar tokenizadores diferentes. Um token também não é uma coordenada do vetor de embedding.')
h('Entrada saída e contexto')
p('A entrada inclui instruções, pergunta, esquema estruturado, documentos e dados da API quando enviados ao modelo. A saída contém o texto ou JSON gerado. Para um modelo causal, o orçamento de contexto precisa acomodar a entrada e a continuação gerada. A configuração atual usa num_ctx=4096; esse número não significa 4.096 palavras nem 4.096 tokens disponíveis somente para a resposta.')
table(['Parâmetro atual','Interpretação'],[
('num_ctx = 4096','Janela configurada por chamada. É necessário contabilizar prompts, dados e espaço para saída.'),
('num_predict = 400','Limite de saída do avaliador/reformulador. Não é a contagem efetivamente consumida; atingir o limite pode truncar JSON.'),
('temperature = 0 ou 0,2','Controle de amostragem do modelo. Não corresponde a graus Celsius nem garante correção factual.'),
('keep_alive = 15m','Retenção solicitada do modelo carregado; não é memória da conversa.')],[5,12])
h('Exemplo de orçamento apenas ilustrativo')
p('Exemplo didático: 4.096 − 3.000 tokens de entrada − 400 de saída = 696 de margem. A entrada deve incluir os marcadores de chat. O protótipo ainda não gerencia esse orçamento automaticamente; esses valores não são medições do projeto.')
h('Por que várias etapas ficam mais lentas')
p('Uma resposta documental comum pode usar uma chamada de avaliação e outra de geração. A reformulação adiciona uma chamada e uma nova avaliação. Cada chamada reprocessa sua própria entrada. Já respostas meteorológicas diretas e a comparação numérica implementada não geram tokens de saída pelo LLM, embora a comparação possa ter passado antes pelo avaliador documental.')
p('Ollama expõe prompt_eval_count, eval_count e durações para medir entrada, saída e tempo. Tokens por segundo podem ser calculados como eval_count dividido por eval_duration em segundos. O uso local não cria cobrança de tokens por um provedor externo, mas consome CPU/GPU, memória, energia e tempo. A aplicação ainda precisa registrar essas métricas por etapa [R6].')

page('7 Exemplo completo de análise de plantio')
p('Entrada: “Quais as condições climáticas hoje? Posso plantar soja?”. O exemplo abaixo corresponde aos dados simulados do projeto e demonstra o encadeamento da informação, não uma orientação para plantio real.')
table(['Origem','Dados relevantes','Interpretação implementada'],[
('API meteorológica','Ar 28 °C; solo 25 °C; chuva 12 mm; vento 8 km/h','Apresenta os valores com unidade, horário e origem [API].'),
('Plano do talhão','Umidade volumétrica 28%; sem encharcamento; trafegável','Verifica preenchimento e atualidade; compara a umidade com a regra fictícia.'),
('Calendário de teste','01/10 a 30/11/2026; consulta em 13/10/2026','Data dentro da janela fictícia; isso não comprova ZARC aplicável.'),
('Documento recuperado','Faixa de 20 a 30 °C para semeadura da soja','Extrai limites do trecho aprovado e faz comparação numérica.')],[3.8,6.4,6.8])
h('Como a execução escolhe a resposta')
p('O grafo consulta a API e valida a medição. Depois consulta o plano da soja e verifica seus vínculos e restrições. Se houver pendências ou condições desfavoráveis nas regras fictícias, encerra a análise combinada com os motivos, sem usar a temperatura para ignorá-los.')
p('Se o plano passar, a busca documental procura condições de temperatura para semeadura. Quando encontra uma frase no formato reconhecido, o código extrai os limites numéricos; não usa 20 e 30 como constantes agronômicas fixas. A comparação informa se as medições estão numericamente dentro ou fora do intervalo.')
h('O ponto que exige cuidado')
p('O trecho recuperado informa a faixa, mas não identifica explicitamente ar ou solo. O programa assinala essa ambiguidade. A página seguinte ou uma passagem complementar pode esclarecer o requisito, mas a recuperação atual não garante essa complementação. Também não valida se a cultivar fictícia realmente é apropriada ao local.')
p('Uma conclusão fiel ao estágio atual é: “O cenário passou nas regras simuladas de solo e calendário. As temperaturas estão numericamente no intervalo recuperado, mas a interpretação desse requisito precisa ser confirmada. Não há evidência suficiente para transformar o resultado em recomendação real de plantio.”')
h('Outras perguntas')
p('“Qual a umidade do solo hoje?” consulta o plano único aplicável e mostra a umidade uma vez. “Qual a temperatura ideal para semear soja?” segue o fluxo documental. “Vai chover amanhã?” não pode ser respondida com a última medição: a API atual não fornece previsão.')

page('8 Verificação qualidade e desempenho')
h('O que foi testado')
p('A última revisão funcional registrou 47 testes Python e quatro testes com JSON Server real. Após reparar a instalação na pasta renomeada, os quatro testes HTTP foram executados novamente. Os testes cobrem decisões do grafo, limite de buscas, fontes inválidas, isolamento de consultas, validade temporal, vínculo do plano, condições desfavoráveis e persistência/recarga do JSON.')
p('Há também registros de execução com Chroma e Ollama reais. Eles mostram que as integrações funcionam em exemplos selecionados. Aprovar testes de software ou percorrer o nó correto não equivale a medir precisão agronômica; um teste anterior seguiu o caminho esperado e, mesmo assim, gerou uma resposta factual incorreta.')
table(['Medição local registrada','Tempo','Condição'],[
('Origem da soja antes da otimização','41,57 s','Diagnóstico com avaliação de 34,09 s.'),
('Origem da soja após otimização','5,58 s','Modelos já carregados; não é comparação controlada.'),
('Umidade do solo após revisão','0,22 s','API e plano, sem geração pelo LLM.'),
('Plantio após revisão da integração','5,62 s','API, plano e recuperação documental.')],[8,2.5,6.5])
p('Os números são observações de execuções distintas, não um benchmark ou promessa de latência. Carregamento inicial, cache, concorrência, tamanho dos textos e divisão CPU/GPU alteram os resultados. Em diagnóstico anterior, o contexto menor aumentou a parcela reportada de GPU de 68% para 78%; isso não é utilização instantânea medida da GPU.')
h('Limites do avaliador e das citações')
p('A avaliação usa o mesmo modelo pequeno que gera respostas. Pydantic verifica estrutura, não a veracidade da conclusão. A conferência de passagens tolera espaços, acentos, caixa e hifenização; ela confirma correspondência textual normalizada, mas não demonstra relevância semântica. Em testes, houve omissão de marcadores de citação e seleção de fonte desnecessária.')
h('Métricas necessárias para uma avaliação real')
bullets(['Recuperação: proporção de perguntas cujo trecho esperado aparece entre os candidatos, separada por cultura e operação.',
'Resposta: sustentação das afirmações, correção das citações, unidades e identificação do talhão e da data.',
'Decisão: falsos favoráveis, falsas restrições, abstenção adequada e discordância com revisão agronômica.',
'Operação: tempo até o primeiro conteúdo útil, latência p50/p95, erros HTTP e consumo de tokens por etapa.'])

page('9 Dados necessários para uma decisão confiável')
p('Um chatbot que apoia decisões precisa de dados corretos, critérios aplicáveis e avaliação independente. Acrescentar campos fictícios ajuda a exercitar o código, mas não substitui coleta, calibração ou validação técnica. Os requisitos abaixo são propostas de evolução, não recursos já disponíveis.')
table(['Grupo de dados','Conteúdo necessário','Controle de qualidade'],[
('Identificação espacial','Fazenda, talhão, coordenadas, área, município e histórico de uso','Vincular sensores e recomendações ao talhão correto; evitar médias sem representatividade.'),
('Solo e água','Textura, profundidade, umidade volumétrica, método, capacidade de campo, drenagem e condição operacional','Calibração, unidade e profundidade de leitura; limiares definidos por especialista para o contexto.'),
('Meteorologia','Séries de chuva, temperatura, vento e previsão quando necessária','Origem, horário de coleta/emissão, horizonte, resolução espacial, lacunas e incerteza.'),
('Cultura e cultivar','Identificação real, grupo/ciclo, sementes e características pertinentes','Cadastro verificável, adequação regional e validade das informações.'),
('Calendário aplicável','Regras oficiais vigentes e suas condições de aplicação','Compatibilizar cultura, município, solo, ciclo/grupo, safra e versão da fonte.'),
('Manejo e operação','Sistema de cultivo, resíduos, irrigação, condição do terreno e objetivo da operação','Critérios revisados por responsável técnico; não inferir apenas de uma medição meteorológica.')],[3.7,6.7,6.6])
h('Zoneamento e atualidade das fontes')
p('O ZARC é um instrumento de gestão de risco climático. Sua aplicação exige consultar os resultados e atos pertinentes ao caso; uma janela genérica de calendário não o substitui. A futura integração deve registrar origem, safra/versão e parâmetros utilizados, em vez de tratar o nome do município como condição suficiente [R7].')
h('Conhecimento documental adequado')
p('O corpus deve separar manuais agronômicos de relatórios técnicos sobre IA, registrar edição, data, cultura, região e condições de validade. Trechos que dependem de tabelas, figuras ou páginas vizinhas precisam manter essas relações. Regras aplicáveis ao solo não podem ser substituídas por medições do ar apenas porque os números são semelhantes.')
p('A política de dados ausentes deve ser explícita: perguntar ao usuário, buscar outra fonte permitida ou abster-se. “Não foi encontrado” não significa zero, ausência de risco ou inexistência da informação em todo o acervo.')

page('10 Arquitetura proposta para um aplicativo')
p('A próxima aplicação pode conservar o núcleo Python e os grafos, mas deve separar interface, serviço de chat e serviços de dados. Esta seção descreve uma proposta; esses componentes não foram implementados no protótipo atual.')
h('Interface e backend de chat')
p('Uma interface web ou móvel permitiria selecionar fazenda, talhão, cultura e operação, apresentar a última atualização dos dados e receber a resposta com fontes. Um backend Python, por exemplo com FastAPI, exporia POST /chat e uma transmissão de eventos por SSE. A interface não deve acessar Chroma ou executar modelos diretamente.')
p('O backend autenticaria o usuário, verificaria autorização para o talhão, atribuiria um identificador à consulta e executaria o LangGraph. O estado da conversa poderia ser persistido por sessão; memória deve guardar identidade e proveniência sem reutilizar automaticamente medições vencidas. Mensagens de progresso não precisam expor detalhes internos irrelevantes ao usuário.')
h('Serviços de dados e persistência')
p('JSON Server seria substituído por APIs reais de estações, sensores, cadastro agrícola e fontes oficiais. Um banco transacional poderia manter fazendas, talhões, permissões e auditoria; séries temporais armazenariam medições. Chroma continuaria como índice documental, com versões e rotina de atualização. Cada armazenamento possui uma finalidade diferente.')
h('Separar regras e linguagem')
p('A camada de decisão deve executar critérios agronômicos explícitos e versionados, aprovados por especialistas. O modelo pode interpretar a pergunta, apoiar a recuperação e explicar o resultado. Ele não deve transformar lacunas em fatos nem criar um limiar operacional sem fonte. Regras conflitantes ou evidências insuficientes devem produzir esclarecimento ou encaminhamento para revisão humana.')
h('Segurança e observabilidade')
bullets(['Autenticação e autorização por fazenda/talhão, TLS, gestão de segredos e proibição de acesso cruzado entre usuários.',
'Permissão de leitura para o chatbot; escritas em cadastro ou comandos sobre equipamentos exigiriam fluxos próprios e autorização.',
'Validação de URLs e ferramentas permitidas; documentos e respostas de API tratados como dados, não instruções.',
'Logs com versão do modelo, regras, fontes, horários e motivos da decisão; proteção de informações pessoais e comerciais.',
'Timeouts, controle de concorrência, limites de consumo, recuperação de falhas e verificação de dados vencidos.'])
p('A API fictícia atual permite operações de alteração de dados e não possui autenticação. Está vinculada ao loopback para testes. Ela não deve ser publicada na internet como backend de produção.')

page('11 Plano de evolução e critérios de aceitação')
h('Etapa 1 Consolidar a engenharia da informação')
p('Corrigir a atualização incremental da ingestão, fixar dependências Python, versionar o corpus e registrar métricas por chamada. Acrescentar recuperação de passagens adjacentes, tratamento de tabelas e seleção por assunto. O critério de aceitação é demonstrar que um documento alterado atualiza corretamente suas evidências e que falha de lote não termina como sucesso.')
h('Etapa 2 Validar o domínio agronômico')
p('Definir com um profissional agronômico quais decisões o produto poderá apoiar. Transformar cada decisão em requisitos de dados, regras, exceções e fontes. Criar uma base de perguntas e casos com resultado esperado, incluindo situações desfavoráveis, limites, dados conflitantes e ausência de dados. Separar casos de desenvolvimento dos casos reservados para avaliação.')
p('Um resultado “favorável” deve exigir evidências aplicáveis e atuais para todos os critérios obrigatórios do escopo. Se um critério não puder ser avaliado, a saída será “dados insuficientes” ou “revisão necessária”. A aprovação final deve incluir revisão agronômica; um número de testes verdes não substitui essa revisão.')
h('Etapa 3 Conectar dados reais e construir a aplicação')
p('Integrar uma fazenda piloto e dados medidos com origem e unidade verificáveis. Substituir a janela fictícia por consulta aplicável e versionada. Implementar backend de chat, identidade do usuário, seleção de talhão, histórico auditável e interface com conclusão primeiro, evidências depois e pendências explícitas.')
h('Etapa 4 Realizar um piloto supervisionado')
p('Comparar o sistema com avaliações independentes, registrar divergências e revisar regras. Medir latência em carga controlada, além de erros de decisão. Ampliar culturas e operações somente após validação de cada escopo; o sucesso do exemplo de soja não valida automaticamente irrigação, pulverização, outras regiões ou outros cultivos.')
table(['Critério de liberação','Evidência esperada'],[
('Dados e proveniência','Nenhuma conclusão usa silenciosamente talhão errado, informação vencida ou fonte sem identificação.'),
('Regras','Condições operacionais justificadas e aprovadas, sem limiares fictícios na execução real.'),
('Respostas','Afirmações sustentadas, citações verificáveis e abstenção correta nos casos definidos como críticos.'),
('Operação','Metas de latência e disponibilidade acordadas e medidas; falhas recuperáveis e trilha de auditoria.')],[4.4,12.6])
p('A equipe técnica e o responsável agronômico devem definir as metas antes do piloto. Ainda não há evidência para declarar uma taxa de acerto ou confiabilidade de campo.')

page('12 Referências e reprodução')
h('Fontes técnicas externas')
refs=[
('R1','LangGraph overview','https://docs.langchain.com/oss/python/langgraph/overview'),
('R2','Retrieval na documentação LangChain','https://docs.langchain.com/oss/python/deepagents/retrieval'),
('R3','JSON Server versão 0.17.4','https://github.com/typicode/json-server/tree/v0.17.4'),
('R4','Introdução ao Chroma','https://docs.trychroma.com/docs/overview/introduction'),
('R5','Tokenization algorithms da Hugging Face','https://huggingface.co/docs/transformers/main/en/tokenizer_summary'),
('R6','Ollama API chat e métricas de uso','https://docs.ollama.com/api/chat'),
('R7','Ministério da Agricultura e Pecuária e ZARC','https://www.gov.br/agricultura/pt-br/assuntos/riscos-seguro/programa-nacional-de-zoneamento-agricola-de-risco-climatico/zoneamento-agricola')]
for num,title,url in refs:
    para=p(f'[{num}] {title}. ')
    hyperlink=OxmlElement('w:hyperlink')
    rel=doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'1F4E79'); props.append(color); run.append(props)
    text=OxmlElement('w:t'); text.text='Abrir referência'; run.append(text); hyperlink.append(run); para._p.append(hyperlink)
p('Referências externas consultadas em 13 de outubro de 2026. As descrições da implementação se referem ao código e aos registros locais do projeto, não a todos os recursos possíveis das bibliotecas.')
h('Evidências locais e configuração')
p('Código principal: src/assistant_graph.py, src/rag_graph.py, src/rag_engine.py, src/weather_api.py, src/planting_context.py e src/ingest.py. API: api/server.cjs, api/seed.cjs, api/db.json e api/pnpm-workspace.yaml. Caminho do banco: config/chroma_path.txt, com precedência de RAG_CHROMA_DIR.')
p('Registros: tests/resultados_langgraph_antes_otimizacao.json, tests/resultados_langgraph_aquecido.json e tests/resultados_correcao_integracao.json. Os registros são amostras históricas; as medições fictícias atuais podem ter sido renovadas pelo seed.')
h('Comandos na raiz da pasta chatbot')
for line in [r'.\iniciar_api.bat',r'.\iniciar_chatbot.bat',r'.\parar_api.bat',r'.\renovar_dados_demo.bat',r'.\.venv312\Scripts\python.exe -m unittest discover -s tests -p "test_*.py"',r'.\.venv312\Scripts\python.exe -m tests.verificar_api_http']:
    para=p(line)
    for run in para.runs: run.font.name='Consolas'; run.font.size=Pt(9)
p('A renovação da simulação substitui os valores editados pelos padrões do seed. Os quatro testes HTTP usam banco temporário; os testes unitários não exigem uma estação meteorológica. Para testes reais do chatbot, Ollama, base vetorial e API precisam estar disponíveis.')

path=OUT/'Relatorio_tecnico_chatbot_agronomico.docx'
doc.save(path)
print(path)
